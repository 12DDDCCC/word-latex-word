"""段落提取模块

提取单个段落的完整信息，包括 run 级别处理、Unicode→LaTeX 转换、化学式检测等。
"""

import re
import sys
from pathlib import Path
from docx.oxml.ns import qn

SKILL_DIR = Path(__file__).resolve().parent.parent
for _path in (SKILL_DIR, SKILL_DIR / 'citation-extract'):
    if str(_path) not in sys.path:
        sys.path.insert(0, str(_path))
from word_link_citations import resolve_citation_key


# Unicode 特殊字符 → LaTeX
UNICODE_TO_LATEX = {
    '₂': '$_{2}$', '₃': '$_{3}$', '₄': '$_{4}$', '₁': '$_{1}$',
    '⁰': '$^{0}$', '²': '$^{2}$', '³': '$^{3}$',
    'λ': '$\\lambda$', 'δ': '$\\delta$', 'α': '$\\alpha$',
    'β': '$\\beta$', 'γ': '$\\gamma$', 'μ': '$\\mu$',
    'σ': '$\\sigma$', 'θ': '$\\theta$', 'φ': '$\\phi$',
    '°': '$^{\\circ}$',
    '×': '$\\times$', '≤': '$\\leq$', '≥': '$\\geq$',
    '±': '$\\pm$', '≈': '$\\approx$',
}

# 公式内 Unicode → LaTeX（不加 $）
UNICODE_TO_LATEX_MATH = {
    '₂': '_{2}', '₃': '_{3}', '₁': '_{1}',
    '⁰': '^{0}', '²': '^{2}', '³': '^{3}',
    'λ': '\\lambda', 'δ': '\\delta',
    '×': '\\times', '°': '\\circ',
}


def _unicode_to_latex(text):
    """将文本中的 Unicode 特殊字符转换为 LaTeX 表示"""
    result = text
    for uc, ltx in UNICODE_TO_LATEX.items():
        result = result.replace(uc, ltx)
    return _restore_identifier_subscripts(result)


def _restore_identifier_subscripts(latex):
    """Do not keep chemical-style subscripts inside dataset/model identifiers."""
    marker = r'(?:\$_\{(?P<m1>\d{1,2})\}\$|\\textsubscript\{(?P<m2>\d{1,2})\}|_\{(?P<m3>\d{1,2})\})'
    token_re = re.compile(r'[A-Za-z][A-Za-z0-9_]*' + marker + r'[A-Za-z_][A-Za-z0-9_]*')

    def repl(m):
        token = m.group(0)
        return re.sub(marker, lambda sm: sm.group('m1') or sm.group('m2') or sm.group('m3'), token)

    return token_re.sub(repl, latex)


def _detect_chemical_formula(latex):
    """检测并转换化学式中的下标

    例如: CO2 → CO$_{2}$, XCO2 → XCO$_{2}$
    只对已知化学元素前缀加下标，非化学式保持原样。
    """
    # 化学式识别：大写字母开头+连续字母+数字 → 下标
    # 条件：数字后面不能紧跟字母或下划线（排除 XCO2quality_flag 等）
    chem_pattern = re.compile(r'(?<![0-9A-Za-z_])([A-Z][A-Za-z]*?)(\d{1,2})(?=[^0-9a-zA-Z_]|$)')
    # 已知化学元素前缀（允许加下标的）
    _CHEM_PREFIXES = {
        'CO', 'XCO', 'H', 'O', 'N', 'C', 'S', 'P', 'Fe', 'Ca', 'Na',
        'K', 'Mg', 'Al', 'Si', 'Cl', 'Ti', 'Mn', 'Zn', 'Cu', 'Pb',
        'NO', 'SO', 'CH', 'NH', 'OH', 'HO', 'NaCl', 'H2',
    }
    # 明确排除的非化学式前缀
    _CHEM_EXCLUDE = {
        'OCO',    # OCO-2/OCO-3 卫星名
        'GCAS',   # GCASv2 模型名
        'GOSAT',  # GOSAT 卫星名
        'TCCON',  # TCCON 网络
        'MODIS',  # MODIS 传感器
        'MOPITT', # MOPITT 传感器
        'CAMS',   # CAMS 模型
        'CMIP',   # CMIP6 项目
        'Net',    # Net CO2 flux
        'GRACED', # GRACED 数据集
        'GFED',   # GFED 数据集
        'BEPS',   # BEPS 模型
        'FLUX',   # FLUX-Site
    }
    chem_matches = []
    for m in chem_pattern.finditer(latex):
        prefix = m.group(1)
        num = m.group(2)
        # 检查是否匹配排除列表（含前缀匹配，如 GCASv 中的 GCAS）
        excluded = False
        for exc in _CHEM_EXCLUDE:
            if prefix.startswith(exc) or exc.startswith(prefix):
                excluded = True
                break
        if excluded:
            continue
        # 只对已知化学元素前缀加下标
        if prefix not in _CHEM_PREFIXES:
            continue
        chem_matches.append((m.start(), m.end(), prefix, num))
    # 从后向前替换，避免偏移问题
    for start, end, prefix, num in reversed(chem_matches):
        latex = latex[:start + len(prefix)] + f'$_{{{num}}}$' + latex[end:]
    return _restore_identifier_subscripts(latex)


def _extract_run(run, color_map):
    """提取单个 run 的完整格式信息"""
    from docx.oxml.ns import qn

    rf = run.font
    text = run.text or ''

    # 颜色
    color_rgb = None
    color_role = None
    if rf.color and rf.color.rgb:
        color_rgb = str(rf.color.rgb)
        color_role = color_map.get(color_rgb, None)

    # 判断引用编号：红色 EE0000 或 themeColor=hyperlink
    is_cite = (color_rgb == 'EE0000')
    if not is_cite:
        # 检查XML中的themeColor属性（Word引用常用hyperlink主题色）
        rPr = run._element.find(qn('w:rPr'))
        if rPr is not None:
            color_el = rPr.find(qn('w:color'))
            if color_el is not None:
                theme = color_el.get(qn('w:themeColor'))
                if theme in ('hyperlink', 'accent1', 'accent2'):
                    is_cite = True

    # LaTeX 转换
    latex = text
    if is_cite:
        nums = re.findall(r'\d+', text)
        if nums:
            keys = ','.join([f'{n}' for n in nums])
            latex = f'\\citep{{{keys}}}'
        else:
            latex = ''
    else:
        # 1) Unicode → LaTeX（先转换，用占位符保护避免后续转义破坏）
        #    占位符格式: \x00UI\x00 (I=字母索引, 无数字, 不会被化学式正则匹配)
        placeholders = {}
        ph_chars = 'ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz'
        ph_idx = 0
        for uc, ltx in UNICODE_TO_LATEX.items():
            c1 = ph_chars[ph_idx % len(ph_chars)]
            c2 = ph_chars[ph_idx // len(ph_chars) % len(ph_chars)]
            ph = f'\x00{c1}{c2}\x00'
            placeholders[ph] = ltx
            ph_idx += 1
            latex = latex.replace(uc, ph)
        # 2) 化学式识别
        latex = _detect_chemical_formula(latex)
        # 3) LaTeX 特殊字符转义（占位符不会被破坏）
        for old, new in [('\\', '\\textbackslash{}'), ('&', '\\&'), ('%', '\\%'),
                         ('#', '\\#'), ('{', '\\{'), ('}', '\\}'),
                         ('_', '\\_'), ('~', '\\textasciitilde{}')]:
            latex = latex.replace(old, new)
        # 4) 恢复占位符
        for ph, ltx in placeholders.items():
            latex = latex.replace(ph, ltx)

    # 格式包裹
    if rf.bold and rf.italic:
        latex = f'\\textbf{{\\textit{{{latex}}}}}'
    elif rf.bold:
        latex = f'\\textbf{{{latex}}}'
    elif rf.italic:
        latex = f'\\textit{{{latex}}}'
    if rf.superscript:
        latex = f'\\textsuperscript{{{latex}}}'
    if rf.subscript:
        latex = f'\\textsubscript{{{latex}}}'
    if color_role and color_role != 'cite' and not is_cite:
        latex = f'{color_role}{{{latex}}}'
    if rf.underline:
        latex = f'\\underline{{{latex}}}'
    if rf.strike:
        latex = f'\\sout{{{latex}}}'  # 需要 soul 包

    return {
        'type': 'text',
        'text': text,
        'bold': rf.bold,
        'italic': rf.italic,
        'underline': rf.underline,
        'strike': rf.strike,
        'superscript': rf.superscript,
        'subscript': rf.subscript,
        'font_name': rf.name,
        'size_pt': _emu_to_pt(rf.size),
        'color_rgb': color_rgb,
        'color_role': color_role,
        'is_cite': is_cite,
        'latex': latex,
    }


def _emu_to_pt(emu):
    """尺寸转换: EMU → pt (1pt = 12700 EMU)"""
    if emu is None:
        return None
    return round(emu / 12700, 1)


def _normalize_adjacent_cite_runs(latex_parts):
    """Merge citation runs split by Word without inventing comma-separated keys."""
    normalized = []
    i = 0
    cite_pat = re.compile(r'^\\citep\{(\d+)\}$')
    while i < len(latex_parts):
        part = latex_parts[i]
        m = cite_pat.match(part or '')
        if not m:
            normalized.append(part)
            i += 1
            continue
        digits = [m.group(1)]
        j = i + 1
        while j < len(latex_parts):
            nm = cite_pat.match(latex_parts[j] or '')
            if not nm:
                break
            digits.append(nm.group(1))
            j += 1
        if len(digits) == 1:
            normalized.append(part)
        elif all(len(d) == 1 for d in digits):
            normalized.append('\\citep{' + ''.join(digits) + '}')
        else:
            normalized.append('\\citep{' + ','.join(digits) + '}')
        i = j
    return normalized


def _run_text_from_elem(run_elem):
    return ''.join(t.text or '' for t in run_elem.findall('.//' + qn('w:t')))


def _field_char_type(run_elem):
    fld = run_elem.find(qn('w:fldChar'))
    return fld.get(qn('w:fldCharType')) if fld is not None else None


def _consume_citation_field(children, start_idx, citation_resolver):
    instruction = ''
    display = ''
    separated = False
    i = start_idx
    while i < len(children):
        child = children[i]
        if child.tag.split('}')[-1] != 'r':
            break
        typ = _field_char_type(child)
        instr = child.find(qn('w:instrText'))
        if typ == 'separate':
            separated = True
        elif typ == 'end':
            i += 1
            break
        elif instr is not None:
            instruction += instr.text or ''
        elif separated:
            display += _run_text_from_elem(child)
        i += 1
    m = re.search(r'HYPERLINK\s+\\l\s+"([^"]+)"', instruction, re.I)
    target = m.group(1) if m else ''
    key = resolve_citation_key(target, display, citation_resolver)
    return i, {'type': 'cite', 'key': key, 'display': display, 'target': target} if key else None


def _cite_latex(keys):
    return '\\citep{' + ','.join(str(k) for k in keys if k) + '}'


def _merge_link_citation_tokens(tokens):
    merged = []
    i = 0
    while i < len(tokens):
        token = tokens[i]
        if not (isinstance(token, str) and token.endswith('(')):
            if isinstance(token, dict) and token.get('type') == 'cite':
                merged.append(_cite_latex([token.get('key')]))
            else:
                merged.append(token)
            i += 1
            continue

        keys = []
        j = i + 1
        if j < len(tokens) and isinstance(tokens[j], dict) and tokens[j].get('type') == 'cite':
            keys.append(tokens[j].get('key'))
            j += 1
            while (
                j + 1 < len(tokens)
                and isinstance(tokens[j], str)
                and re.fullmatch(r'\s*[;,]\s*', tokens[j] or '')
                and isinstance(tokens[j + 1], dict)
                and tokens[j + 1].get('type') == 'cite'
            ):
                keys.append(tokens[j + 1].get('key'))
                j += 2
            if j < len(tokens) and isinstance(tokens[j], str) and tokens[j].startswith(')'):
                prefix = token[:-1]
                suffix = tokens[j][1:]
                if prefix:
                    merged.append(prefix)
                merged.append(_cite_latex(keys))
                if suffix:
                    merged.append(suffix)
                i = j + 1
                continue

        merged.append(token)
        i += 1
    return [str(part) for part in merged]


def _merge_parenthetical_cite_latex(latex):
    def repl(match):
        keys = []
        for group in re.findall(r'\\citep\{([^}]+)\}', match.group(1)):
            keys.extend(k.strip() for k in group.split(',') if k.strip())
        return _cite_latex(keys) if keys else match.group(0)

    return re.sub(
        r'\((\\citep\{[^}]+\}(?:\s*[;,]\s*\\citep\{[^}]+\})+)\)',
        repl,
        latex,
    )


def extract_paragraph(para, pi, heading_level_func, alignment_name_func, omml_to_latex_func,
                      color_map, citation_resolver=None):
    """提取单个段落的完整信息

    Args:
        para: python-docx Paragraph 对象
        pi: 段落索引
        heading_level_func: 标题级别推断函数
        alignment_name_func: 对齐方式转换函数
        omml_to_latex_func: OMML→LaTeX 转换函数
        color_map: 颜色映射表

    Returns:
        dict: {
            'para_index': int,
            'style': str,
            'heading_level': int or None,
            'alignment': str or None,
            'first_line_indent_pt': float or None,
            'left_indent_pt': float or None,
            'line_spacing': float or None,
            'space_before_pt': float or None,
            'space_after_pt': float or None,
            'runs': list[dict],  # run级别格式
            'text': str,         # 纯文本
            'latex': str,        # LaTeX 格式文本（含公式）
            'has_formula': bool,
        }
    """
    from docx.oxml.ns import qn

    pf = para.paragraph_format
    style = para.style.name if para.style else 'Normal'
    h_level = heading_level_func(style)

    result = {
        'para_index': pi,
        'style': style,
        'heading_level': h_level,
        'semantic_type': None,  # 后续由 classify_semantic_type() 填充
        'alignment': alignment_name_func(pf.alignment),
        'first_line_indent_pt': _emu_to_pt(pf.first_line_indent),
        'left_indent_pt': _emu_to_pt(pf.left_indent),
        'line_spacing': pf.line_spacing if pf.line_spacing else None,
        'space_before_pt': _emu_to_pt(pf.space_before),
        'space_after_pt': _emu_to_pt(pf.space_after),
        'runs': [],
        'text': para.text,
        'latex': '',
        'has_formula': False,
    }

    # 检查公式
    omaths = para._element.findall('.//' + qn('m:oMath'))
    omath_paras = para._element.findall('.//' + qn('m:oMathPara'))
    if omaths or omath_paras:
        result['has_formula'] = True

    # 提取 run 级别格式 + 拼接 LaTeX
    latex_parts = []
    elem = para._element

    # 遍历段落直接子节点，处理 run 和公式交替出现
    children = list(elem)
    child_idx = 0
    while child_idx < len(children):
        child = children[child_idx]
        tag = child.tag.split('}')[-1]

        if tag == 'r':
            if _field_char_type(child) == 'begin':
                next_idx, cite_token = _consume_citation_field(
                    children, child_idx, citation_resolver)
                if cite_token:
                    result['runs'].append({
                        'type': 'citation_link',
                        'text': cite_token['display'],
                        'key': cite_token['key'],
                        'target': cite_token['target'],
                        'latex': _cite_latex([cite_token['key']]),
                    })
                    latex_parts.append(cite_token)
                    child_idx = next_idx
                    continue

            # 文本 run
            run_obj = None
            for run in para.runs:
                if run._element is child:
                    run_obj = run
                    break
            if run_obj is None:
                child_idx += 1
                continue

            run_info = _extract_run(run_obj, color_map)
            result['runs'].append(run_info)
            latex_parts.append(run_info['latex'])

        elif tag == 'oMath':
            # 行内公式
            latex = omml_to_latex_func(child)
            if latex:
                result['runs'].append({
                    'type': 'formula',
                    'latex': latex,
                    'formula_type': 'inline',
                })
                latex_parts.append(f'$ {latex} $')

        elif tag == 'oMathPara':
            # 独立段落公式
            for om in child.findall(qn('m:oMath')):
                latex = omml_to_latex_func(om)
                if latex:
                    result['runs'].append({
                        'type': 'formula',
                        'latex': latex,
                        'formula_type': 'display',
                    })
                    latex_parts.append(f'\\begin{{equation}}\n  {latex}\n\\end{{equation}}')

        elif tag == 'hyperlink':
            # 超链接
            rId = child.get(qn('r:id'))
            url = ''
            if rId:
                try:
                    rel = para.part.rels[rId]
                    url = rel.target_ref
                except Exception:
                    pass
            link_text = ''.join(t.text for t in child.findall('.//' + qn('w:t')) if t.text)
            if link_text and url:
                result['runs'].append({
                    'type': 'hyperlink',
                    'text': link_text,
                    'url': url,
                    'latex': f'\\url{{{url}}}',
                })
                latex_parts.append(f'\\url{{{url}}}')
            elif link_text:
                anchor = child.get(qn('w:anchor'))
                key = resolve_citation_key(anchor, link_text, citation_resolver)
                if key:
                    result['runs'].append({
                        'type': 'citation_link',
                        'text': link_text,
                        'target': anchor,
                        'key': key,
                        'latex': _cite_latex([key]),
                    })
                    latex_parts.append({
                        'type': 'cite',
                        'key': key,
                        'display': link_text,
                        'target': anchor,
                    })
                else:
                    result['runs'].append({
                        'type': 'hyperlink',
                        'text': link_text,
                        'url': '',
                        'latex': link_text,
                    })
                    latex_parts.append(link_text)

        child_idx += 1

    latex_parts = _merge_link_citation_tokens(latex_parts)
    latex_parts = _normalize_adjacent_cite_runs(latex_parts)
    result['latex'] = _merge_parenthetical_cite_latex(''.join(latex_parts))
    return result
